"""Independent validation of the saved CSV package against untouched sources."""
import collections, csv, datetime, hashlib, json, pathlib, re, warnings
from decimal import Decimal
import openpyxl

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
ROOT=pathlib.Path.cwd(); OUT=ROOT/'_data/data-for-import'
manifest=json.loads((OUT/'manifest.json').read_text(encoding='utf-8'))
checks=[]; data={}; headers={}
def test(name,condition,details=''):
    checks.append({'check':name,'result':'pass' if condition else 'FAIL','detail':details})
def dec(s):return Decimal(s) if s!='' else None
for f in manifest['files']:
    p=OUT/f['filename']; content=p.read_bytes()
    test('SHA-256 '+p.name,hashlib.sha256(content).hexdigest()==f['sha256'])
    test('UTF-8 BOM '+p.name,content.startswith(b'\xef\xbb\xbf'))
    with p.open(encoding='utf-8-sig',newline='') as stream:
        reader=csv.DictReader(stream); records=list(reader); headers[p.name]=reader.fieldnames
    data[p.name]=records
    test('row count '+p.name,len(records)==f['row_count'])
    test('rectangular CSV '+p.name,all(None not in x and all(v is not None for v in x.values()) for x in records))
    test('unique headers '+p.name,len(set(reader.fieldnames))==len(reader.fieldnames))
    for field in reader.fieldnames:
        if field.endswith('_json') or field in ['merged_ranges']:
            try:
                for record in records:
                    if record[field]!='':json.loads(record[field])
                test('JSON '+p.name+'.'+field,True)
            except Exception:test('JSON '+p.name+'.'+field,False)

keys={
 'sources.csv':['source_id'],'source_rows.csv':['source_row_id'],
 'source_sheets.csv':['source_id','sheet'],'source_tables.csv':['source_id','sheet','table'],
 'source_cells.csv':['source_row_id','cell'],'source_media.csv':['source_id','zip_path'],
 'categories.csv':['category_id'],'collections.csv':['collection_id'],'items.csv':['item_id'],
 'item_sources.csv':['item_id','source_row_id','role'],'identifiers.csv':['identifier_id'],
 'packaging.csv':['package_id'],'locations.csv':['location_id'],
 'quantity_observations.csv':['observation_id'],'opening_balances.csv':['opening_id'],
 'duplicate_identifiers.csv':['source_id','sheet','identifier'],
 'quantity_reconciliation.csv':['item_id','location_id'],
 'organizations.csv':['organization_id'],'item_organizations.csv':['item_id','organization_id','source_row_id'],
 'purchase_observations.csv':['observation_id'],'purchase_orders.csv':['order_id'],
 'purchase_lines.csv':['line_id'],'purchase_line_sources.csv':['line_id','observation_id'],
 'purchase_cost_components.csv':['source_row_id'],'purchased_product_summaries.csv':['source_row_id'],
 'external_listings.csv':['listing_id'],'shipping_details.csv':['item_id'],
 'supplier_product_observations.csv':['model','observed_on'],'item_values.csv':['value_id'],
 'issues.csv':['issue_id'],'validation.csv':['check']}
for name,fields in keys.items():
    test('primary key '+name,len({tuple(r[k] for k in fields) for r in data[name]})==len(data[name]) and all(all(r[k]!='' for k in fields) for r in data[name]))
idsets={k:{r[k] for r in data[t]} for k,t in [('source_id','sources.csv'),('source_row_id','source_rows.csv'),('category_id','categories.csv'),('collection_id','collections.csv'),('item_id','items.csv'),('location_id','locations.csv'),('organization_id','organizations.csv'),('order_id','purchase_orders.csv'),('line_id','purchase_lines.csv')]}
quantity_ids={r['observation_id'] for r in data['quantity_observations.csv']}
purchase_ids={r['observation_id'] for r in data['purchase_observations.csv']}
for name,records in data.items():
    for key,valid in idsets.items():
        if key in headers[name]:test('foreign key '+name+'.'+key,all(not r[key] or r[key] in valid for r in records))
test('opening observation references',all(r['observation_id'] in quantity_ids for r in data['opening_balances.csv']))
test('purchase source references',all(r['observation_id'] in purchase_ids for r in data['purchase_line_sources.csv']))
test('purchase summary references',all(not r['matched_observation_id'] or r['matched_observation_id'] in purchase_ids for r in data['purchased_product_summaries.csv']))
test('purchase component candidate references',all(not r['candidate_line_id'] or r['candidate_line_id'] in idsets['line_id'] for r in data['purchase_cost_components.csv']))
test('every purchase occurrence retained',collections.Counter(r['observation_id'] for r in data['purchase_line_sources.csv'])==collections.Counter(r['observation_id'] for r in data['purchase_observations.csv']))
selected=collections.Counter(r['line_id'] for r in data['purchase_line_sources.csv'] if r['selected']=='true')
test('one preferred source per line',all(selected[r['line_id']]==1 for r in data['purchase_lines.csv']))

books={}; cached={}; coordinate_keys=set()
for source in data['sources.csv']:
    p=ROOT/'_data/current-inventory'/source['filename'] if source['source_kind']=='workbook' else ROOT/source['filename']
    test('source unchanged '+source['source_id'],hashlib.sha256(p.read_bytes()).hexdigest()==source['sha256'])
    if source['source_kind']!='workbook':continue
    k=source['source_id'];books[k]=openpyxl.load_workbook(p,data_only=False);cached[k]=openpyxl.load_workbook(p,data_only=True)
    for s in books[k]:
        for row in s:
            for c in row:
                if c.value is not None:coordinate_keys.add((f'{k}:{s.title}:{c.row}',c.coordinate))
test('all and only populated workbook cells preserved',coordinate_keys=={(r['source_row_id'],r['cell']) for r in data['source_cells.csv']})
test('all sheets covered', {(k,s.title) for k,w in books.items() for s in w}=={(r['source_id'],r['sheet']) for r in data['source_sheets.csv']})
test('source evidence retains all formula cells',sum(c.data_type=='f' for w in books.values() for s in w for row in s for c in row)==sum(r['excel_type']=='f' for r in data['source_cells.csv']))
test('formula cache inventory accounted',sum(r['cache_state'] in ['value','empty_string','error','missing'] for r in data['source_cells.csv'])==sum(r['excel_type']=='f' for r in data['source_cells.csv']))

# Verify exported literal and cached values against the original or an explicit
# restricted crosswalk. No restricted strings are printed or copied into reports.
cross=json.loads((ROOT/'_secrets/data-import-review/restricted-source-text.json').read_text(encoding='utf-8'))
allowed=collections.defaultdict(set)
for r in cross:allowed[r['original']].add(r['sanitized'])
def canonical(v):
    return v.isoformat(timespec='seconds') if isinstance(v,datetime.datetime) else v.isoformat() if isinstance(v,datetime.date) else v
badvalues=0
for row in data['source_cells.csv']:
    k,s,_=row['source_row_id'].split(':');coord=row['cell']
    for field,book in [('value_json',books),('cached_value_json',cached)]:
        original=canonical(book[k][s][coord].value);actual=json.loads(row[field])
        if actual!=original and not (isinstance(original,str) and actual in allowed[original]):badvalues+=1
test('all literal/cache values preserved or explicitly sanitized',badvalues==0,str(badvalues)+' mismatches')

openings=data['opening_balances.csv']; obs={x['observation_id']:x for x in data['quantity_observations.csv']}
test('no opening duplicate pairs',len({(r['item_id'],r['location_id']) for r in openings})==len(openings))
test('all openings require production confirmation',all(r['approved_for_production']=='false' and r['test_fixture_only']=='true' and r['effective_posting_date']=='' and r['note']=='starting inventory' for r in openings))
test('opening source role/quantity exact',all(obs[r['observation_id']]['role']=='location_total' and obs[r['observation_id']]['opening_candidate']=='true' and obs[r['observation_id']]['exclusion_reason']=='' and dec(obs[r['observation_id']]['quantity'])==dec(r['quantity_each']) for r in openings))
test('opening units are whole',all(dec(r['quantity_each'])==dec(r['quantity_each']).to_integral_value() for r in openings))
test('no shipping/game/channel/summaries in openings',all(obs[r['observation_id']]['source_row_id'].startswith('main:') for r in openings))
no_duplicate=True;all_literal_or_counted=True
for r in openings:
    row=obs[r['observation_id']];k,s,rr=row['source_row_id'].split(':');rr=int(rr);sku=cached[k][s].cell(rr,1).value
    table=next(iter(books[k][s].tables.values()));minc,minr,maxc,maxr=openpyxl.utils.range_boundaries(table.ref)
    matches=[n for n in range(minr+1,maxr+1) if cached[k][s].cell(n,1).value==sku]
    no_duplicate &=len(matches)==1 and not sku.endswith('_AAAA')
    sourcecell=books[k][s][row['source_cell']]
    if sourcecell.data_type=='f':
        countparts=[cached[k][s].cell(rr,c).value for c in range(4,9)]
        all_literal_or_counted &=any(x is not None and x!='' for x in countparts)
    all_literal_or_counted &=dec(r['quantity_each'])==Decimal(str(cached[k][s][row['source_cell']].value))
test('openings exclude duplicate identifiers/headings',no_duplicate)
test('opening equals source and has actual count inputs',all_literal_or_counted)

mainitems={r['legacy_sku']:r for r in data['items.csv'] if r['legacy_sku']}
source_skus={cached['main']['Catalog'].cell(r,5).value for r in range(2,160) if cached['main']['Catalog'].cell(r,4).value!='AAAA'}
test('main SKU set exactly preserved',set(mainitems)==source_skus)
test('no fabricated Unit Cost or donation asks',all(r['unit_cost']==r['in_person_ask']==r['online_ask']=='' for r in data['items.csv']))
for r in data['items.csv']:
    if r['source_row_id'].startswith('games:'):
        n=int(r['source_row_id'].split(':')[-1]);test('games FMV follows source MSRP '+str(n),dec(r['irs_fmv'])==Decimal(str(cached['games']['Inventory'].cell(n,5).value)))
test('purchase references never post inventory',all(r['inventory_effect']=='none' for name in ['purchase_orders.csv','purchase_lines.csv','external_listings.csv','supplier_product_observations.csv'] for r in data[name]))
test('purchase descriptions do not contain formula errors',all(r['description'] not in ['#REF!','#VALUE!','#N/A'] for r in data['purchase_lines.csv']))
test('all builder checks pass',all(r['result']=='pass' for r in data['validation.csv']))

# Independent source controls and documented exceptions.
g=cached['games']['Inventory'];game_total=sum(Decimal(str(g.cell(r,4).value))*Decimal(str(g.cell(r,5).value)) for r in range(2,28))
test('games source total cents',game_total.quantize(Decimal('.01'))==Decimal('9678.62'))
test('games quantities',sum(g.cell(r,4).value for r in range(2,28))==341)
test('three zero game titles',sum(g.cell(r,4).value==0 for r in range(2,28))==3)
test('two per-100 price basis exceptions',sum(r['price_basis']=='per_100_candidate' for r in data['purchased_product_summaries.csv'])==2)
test('57 items require internal SKU assignment',sum(not r['proposed_sku'] for r in data['items.csv'])==57)
test('all 193 items have source links', {r['item_id'] for r in data['item_sources.csv']}==idsets['item_id'])
test('main latest purchase is May 2024',max(r['source_date'] for r in data['purchase_observations.csv'] if r['source_row_id'].startswith('main:'))=='2024-05-02')

numeric={'bytes','sheet_index','max_row','max_column','nonempty_cells','formula_cells','cached_errors','header_rows','totals_rows','row','column','quantity','quantity_each','quantity_each','amount','unit_cost','irs_fmv','in_person_ask','online_ask','extended_cost','bundle_quantity','priority','channel_quantity','candidate_units_per_listing','inside_length','inside_width','inside_height','outside_length','outside_width','outside_height','estimated_empty_weight','estimated_packed_weight','outside_volume_ft3','summary_first_quantity','location_numeric_sum_for_audit_only','selected_opening_quantity','item_cost','discount','shipping','tax','fee','total','last_quantity_ordered','unit_price_label_amount','times_ordered','matched_price_each','label_to_each_ratio','inside_length_in','inside_width_in','inside_height_in','outside_length_in','outside_width_in','outside_height_in','outside_tolerance_in','bale_quantity','bundle_weight_lb','unit_weight_lb','price_each','price_min_quantity'}
booleans={'hidden_row','opening_candidate','approved_for_production','test_fixture_only','selected'}
dates={'extracted_on','source_as_of','observed_on','observed_as_of','effective_posting_date','source_date','last_order_date','as_of'}
schema={'schema_version':'draft-1','purpose':'CSV staging contract, not approved database schema','encoding':'UTF-8 with BOM','null':'empty CSV field; JSON fields preserve their own null/empty types','decimal_handling':'Parse exact decimal strings; do not coerce IDs or JSON through numeric conversion.','files':{}}
for name,records in data.items():
    cols=[]
    for col in headers[name]:
        typ='json' if col.endswith('_json') or col=='merged_ranges' else 'boolean' if col in booleans else 'decimal' if col in numeric else 'date' if col in dates else 'string'
        cols.append({'name':col,'type':typ,'nullable':col not in keys[name]})
        valid=True
        for row in records:
            val=row[col]
            if not val:continue
            try:
                if typ=='decimal':valid &=Decimal(val).is_finite()
                elif typ=='boolean':valid &=val in ['true','false']
                elif typ=='date':datetime.date.fromisoformat(val)
            except Exception:valid=False
        if typ in ['decimal','boolean','date']:test('typed values '+name+'.'+col,valid)
    refs=[]
    for key,target in [('source_id','sources.csv'),('source_row_id','source_rows.csv'),('category_id','categories.csv'),('collection_id','collections.csv'),('item_id','items.csv'),('location_id','locations.csv'),('organization_id','organizations.csv'),('order_id','purchase_orders.csv'),('line_id','purchase_lines.csv')]:
        if key in headers[name] and name!=target:refs.append({'column':key,'references_file':target,'references_column':key,'blank_allowed':True})
    if name=='opening_balances.csv':refs.append({'column':'observation_id','references_file':'quantity_observations.csv','references_column':'observation_id','blank_allowed':False})
    if name=='purchase_line_sources.csv':refs.append({'column':'observation_id','references_file':'purchase_observations.csv','references_column':'observation_id','blank_allowed':False})
    schema['files'][name]={'primary_key':keys[name],'foreign_keys':refs,'columns':cols,'purpose':next(x['purpose'] for x in manifest['files'] if x['filename']==name)}
(OUT/'schema.json').write_text(json.dumps(schema,indent=2),encoding='utf-8')
repeat_path=OUT/'repeatability.json'
if repeat_path.exists():
    repeat=json.loads(repeat_path.read_text(encoding='utf-8'))
    current={x['filename']:x['sha256'] for x in manifest['files']}
    test('repeatability proof matches this CSV package',repeat['before']==repeat['after']==current)
result={'checks':len(checks),'passed':sum(c['result']=='pass' for c in checks),'failed':[c for c in checks if c['result']!='pass'],'checks_detail':checks,'limits':['No application importer exists.','No physical stock count performed.','Original XLSX formulas not recalculated; arithmetic was independently checked where specified.','Live ULINE observation is limited to the documented sample.']}
(OUT/'verification.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
(OUT/'verification.md').write_text('# Import package verification\n\n'+f'Checks: {result["checks"]}; passed: {result["passed"]}; failed: {len(result["failed"])}.\n\n'+'Validated standard CSV parsing, checksums, keys/joins, all populated workbook cell coverage, unchanged sources, literal and cached-value preservation, opening eligibility, unit/price exceptions and typed values. `schema.json` describes every CSV column type and primary key. See `verification.json` for individual results.\n\n'+'This verifies the extraction package, not production stock accuracy or an end-to-end application import. The sources were not resaved or recalculated.\n',encoding='utf-8')
print(json.dumps({'checks':result['checks'],'passed':result['passed'],'failed':result['failed']},indent=2))
raise SystemExit(bool(result['failed']))
